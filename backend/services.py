# -*- coding: utf-8 -*-
"""业务编排：采集预览、异步填充、审计导出与历史查询。"""

import json
import re
import threading
import traceback
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from werkzeug.datastructures import FileStorage

import ai_rewrite
import excel_writer
import git_source
import models
from models import (
    AUDIT_DIR,
    LOG_DIR,
    SAMPLES_DIR,
    SNAPSHOT_DIR,
    FillCommit,
    FillRow,
    FillRun,
    SessionLocal,
)

_fill_lock = threading.Lock()


class FillBusy(RuntimeError):
    pass


class ValidationError(RuntimeError):
    pass


class RunLogger:
    def __init__(self, run_id, path):
        self.run_id = run_id
        self.path = Path(path)
        self.lines = []

    def _append(self, level, message):
        line = "[{}] [{}] {}".format(
            datetime.now().strftime("%Y-%m-%d %H:%M:%S"), level, message
        )
        self.lines.append(line)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        with self.path.open("a", encoding="utf-8") as handle:
            handle.write(line + "\n")

    def info(self, message):
        self._append("INFO", message)

    def warn(self, message):
        self._append("WARN", message)

    def error(self, message):
        self._append("ERROR", message)

    def exception(self, prefix=""):
        self._append("ERROR", (prefix + "\n" if prefix else "") + traceback.format_exc())

    def text(self):
        return "\n".join(self.lines)


def parse_repos(raw):
    if isinstance(raw, (list, tuple)):
        return [str(item).strip() for item in raw if str(item).strip()]
    text = str(raw or "").strip()
    if not text:
        return []
    if text.startswith("["):
        try:
            return parse_repos(json.loads(text))
        except (ValueError, TypeError):
            pass
    return [
        line.strip()
        for line in text.replace("\r", "\n").split("\n")
        if line.strip()
    ]


def parse_aliases(raw):
    if isinstance(raw, dict):
        return {
            str(field): [str(name) for name in names if str(name).strip()]
            for field, names in raw.items()
        }
    try:
        parsed = json.loads(raw or "{}")
        return parse_aliases(parsed)
    except (ValueError, TypeError):
        return dict(models.DEFAULT_ALIASES)


def parse_bool(raw, default=False):
    if isinstance(raw, bool):
        return raw
    if raw is None:
        return default
    return str(raw).strip().lower() in ("1", "true", "yes", "on")


def parse_keywords(raw):
    if isinstance(raw, (list, tuple)):
        return [str(item).strip() for item in raw if str(item).strip()]
    return [item.strip() for item in re.split(r"[,，\n]", str(raw or "")) if item.strip()]


def load_options():
    configs = models.get_all_configs()
    return {
        "sharepoint_url": configs.get("sharepoint_url", "").strip(),
        "write_mode": (configs.get("write_mode") or "open").strip().lower(),
        "target_sheet": configs.get("target_sheet", "").strip(),
        "target_mode": (configs.get("target_mode") or "existing").strip().lower(),
        "weekly_template_sheet": configs.get("weekly_template_sheet", "").strip(),
        "weekly_sheet_prefix": configs.get("weekly_sheet_prefix", "周报汇总").strip() or "周报汇总",
        "weekly_overwrite": parse_bool(configs.get("weekly_overwrite"), True),
        "header_row": int(configs.get("header_row") or 0),
        "author": configs.get("author", "").strip(),
        "author_display": configs.get("author_display", "").strip(),
        "repos": parse_repos(configs.get("repos")),
        "week_start_day": int(configs.get("week_start_day") or 0),
        "footer_keywords": parse_keywords(configs.get("footer_keywords")),
        "aliases": parse_aliases(configs.get("column_aliases")),
        "date_mode": (configs.get("date_mode") or "last").strip().lower(),
        "aggregate_mode": (configs.get("aggregate_mode") or "daily").strip().lower(),
        "reuse_empty_rows": parse_bool(configs.get("reuse_empty_rows"), True),
        "dry_run": parse_bool(configs.get("dry_run"), True),
        "copy_row_format": parse_bool(configs.get("copy_row_format"), True),
        "insert_before_footer": parse_bool(configs.get("insert_before_footer"), True),
        "ignore_merge_commits": parse_bool(
            configs.get("ignore_merge_commits"), True
        ),
        "max_description_chars": int(configs.get("max_description_chars") or 1000),
        "sample_path": configs.get("sample_path", "").strip(),
        "ai_enabled": parse_bool(configs.get("ai_enabled"), False),
        "ai_base_url": configs.get("ai_base_url", "").strip(),
        "ai_api_key": configs.get("ai_api_key", "").strip(),
        "ai_model": configs.get("ai_model", "").strip(),
        "ai_prompt": configs.get("ai_prompt", "").strip(),
        "ai_timeout": int(configs.get("ai_timeout") or 60),
        "ai_temperature": (
            float(configs["ai_temperature"])
            if configs.get("ai_temperature") not in (None, "")
            else 0.3
        ),
    }


def public_config():
    options = load_options()
    return {
        "sharepoint_url": options["sharepoint_url"],
        "write_mode": options["write_mode"],
        "target_sheet": options["target_sheet"],
        "target_mode": options["target_mode"],
        "weekly_template_sheet": options["weekly_template_sheet"],
        "weekly_sheet_prefix": options["weekly_sheet_prefix"],
        "weekly_overwrite": options["weekly_overwrite"],
        "header_row": options["header_row"],
        "author": options["author"],
        "author_display": options["author_display"],
        "repos": options["repos"],
        "week_start_day": options["week_start_day"],
        "footer_keywords": options["footer_keywords"],
        "column_aliases": options["aliases"],
        "date_mode": options["date_mode"],
        "aggregate_mode": options["aggregate_mode"],
        "reuse_empty_rows": options["reuse_empty_rows"],
        "dry_run": options["dry_run"],
        "copy_row_format": options["copy_row_format"],
        "insert_before_footer": options["insert_before_footer"],
        "ignore_merge_commits": options["ignore_merge_commits"],
        "max_description_chars": options["max_description_chars"],
        "sample_path": options["sample_path"],
        "ai_enabled": options["ai_enabled"],
        "ai_base_url": options["ai_base_url"],
        "ai_api_key": options["ai_api_key"],
        "ai_model": options["ai_model"],
        "ai_prompt": options["ai_prompt"],
        "ai_timeout": options["ai_timeout"],
        "ai_temperature": options["ai_temperature"],
    }


ALLOWED_CONFIG_KEYS = {
    "sharepoint_url",
    "write_mode",
    "target_sheet",
    "target_mode",
    "weekly_template_sheet",
    "weekly_sheet_prefix",
    "weekly_overwrite",
    "header_row",
    "author",
    "author_display",
    "repos",
    "week_start_day",
    "footer_keywords",
    "column_aliases",
    "date_mode",
    "aggregate_mode",
    "reuse_empty_rows",
    "dry_run",
    "copy_row_format",
    "insert_before_footer",
    "ignore_merge_commits",
    "max_description_chars",
    "sample_path",
    "ai_enabled",
    "ai_base_url",
    "ai_api_key",
    "ai_model",
    "ai_prompt",
    "ai_timeout",
    "ai_temperature",
}


def update_config(payload):
    if not isinstance(payload, dict):
        raise ValidationError("配置必须是对象")
    for key, value in payload.items():
        if key not in ALLOWED_CONFIG_KEYS:
            continue
        if key == "repos":
            value = json.dumps(parse_repos(value), ensure_ascii=False)
        elif key == "column_aliases":
            value = json.dumps(parse_aliases(value), ensure_ascii=False)
        elif key in (
            "dry_run",
            "copy_row_format",
            "insert_before_footer",
            "ignore_merge_commits",
            "reuse_empty_rows",
            "ai_enabled",
            "weekly_overwrite",
        ):
            value = "1" if parse_bool(value) else "0"
        elif key == "footer_keywords":
            value = ",".join(parse_keywords(value))
        elif key == "header_row":
            value = str(max(0, int(value or 0)))
        elif key == "week_start_day":
            value = str(min(6, max(0, int(value or 0))))
        elif key == "max_description_chars":
            value = str(max(100, int(value or 1000)))
        elif key == "ai_timeout":
            value = str(max(10, int(value or 60)))
        elif key == "ai_temperature":
            number = 0.3 if value in (None, "") else float(value)
            value = str(min(2.0, max(0.0, number)))
        elif key == "date_mode":
            value = str(value or "last").strip().lower()
            if value not in ("last", "first", "week_start", "range"):
                raise ValidationError("日期写入方式不正确")
        elif key == "target_mode":
            value = str(value or "existing").strip().lower()
            if value not in ("existing", "weekly", "both"):
                raise ValidationError("写入目标必须是 existing、weekly 或 both")
        elif key == "aggregate_mode":
            value = str(value or "daily").strip().lower()
            if value not in ("daily", "ticket"):
                raise ValidationError("汇总粒度必须是 daily 或 ticket")
        elif key == "write_mode":
            value = str(value or "attach").strip().lower()
            if value not in ("attach", "open"):
                raise ValidationError("写入方式必须是 attach 或 open")
        elif key == "sharepoint_url":
            value = str(value or "").strip()
            if value and not value.lower().startswith(("http://", "https://")):
                raise ValidationError("SharePoint 链接必须以 http(s):// 开头")
        else:
            value = str(value or "").strip()
        models.set_config(key, value)
    return public_config()


def _parse_anchor(value):
    if isinstance(value, date):
        return value
    try:
        return datetime.strptime(str(value)[:10], "%Y-%m-%d").date()
    except (ValueError, TypeError):
        return date.today()


def build_preview(
    anchor_value,
    author=None,
    repos=None,
    start_day=None,
    range_start=None,
    range_end=None,
):
    options = load_options()
    if range_start and range_end:
        start = _parse_anchor(range_start)
        end = _parse_anchor(range_end)
        if start > end:
            raise ValidationError("开始日期不能晚于结束日期")
        if (end - start).days > 366:
            raise ValidationError("时间范围不能超过 366 天")
        anchor = start
    else:
        anchor = _parse_anchor(anchor_value)
        start_day = options["week_start_day"] if start_day is None else int(start_day)
        start, end = git_source.week_bounds(anchor, start_day)
    selected_repos = parse_repos(repos) if repos is not None else options["repos"]
    selected_author = options["author"] if author is None else str(author).strip()

    if not selected_repos:
        raise ValidationError("还没有配置 Git 仓库路径")
    missing = [path for path in selected_repos if not Path(path).exists()]
    if missing:
        raise ValidationError("仓库路径不存在: {}".format("；".join(missing)))

    commits = git_source.load_commits(
        selected_repos,
        selected_author,
        start,
        end,
        ignore_merges=options["ignore_merge_commits"],
    )
    if options["target_mode"] == "weekly":
        aggregate_mode = "ticket"
    elif options["target_mode"] == "both":
        aggregate_mode = "daily"
    else:
        aggregate_mode = options["aggregate_mode"]
    rows = git_source.aggregate_commits(
        commits, aggregate_mode, options["max_description_chars"]
    )
    warnings = []
    if options["ai_enabled"]:
        ai_rewrite.rewrite_rows(rows, options, warnings)
    days = (end - start).days + 1
    daily = {}
    for offset in range(days):
        daily[(start + timedelta(days=offset)).toordinal()] = 0
    for commit in commits:
        key = commit.committed_at.date().toordinal()
        daily[key] = daily.get(key, 0) + 1

    return {
        "week_start": start.isoformat(),
        "week_end": end.isoformat(),
        "range_start": start.isoformat(),
        "range_end": end.isoformat(),
        "range_days": days,
        "anchor": anchor.isoformat(),
        "author": selected_author,
        "repos": selected_repos,
        "repo_count": len(selected_repos),
        "commits_count": len(commits),
        "ticket_count": len({row["ticket"] for row in rows}),
        "row_count": len(rows),
        "aggregate_mode": aggregate_mode,
        "target_mode": options["target_mode"],
        "weekly_template_sheet": options["weekly_template_sheet"],
        "weekly_sheet_name": (
            excel_writer.weekly_sheet_name(options["weekly_sheet_prefix"], start)
            if options["target_mode"] != "existing"
            else ""
        ),
        "ai_enabled": options["ai_enabled"],
        "ai_model": options["ai_model"] if options["ai_enabled"] else "",
        "warnings": warnings,
        "daily": [
            {
                "date": (start + timedelta(days=offset)).isoformat(),
                "count": daily.get(
                    (start + timedelta(days=offset)).toordinal(), 0
                ),
            }
            for offset in range(days)
        ],
        "rows": rows,
        "commits": [commit.to_dict() for commit in commits],
    }


def _number(value):
    try:
        text = str(value or "").strip().replace("%", "").replace(",", "")
        return float(text) if text else 0.0
    except ValueError:
        return 0.0


def _enrich_weekly_rows(daily_rows, weekly_rows, options):
    by_ticket = {}
    for row in daily_rows:
        by_ticket.setdefault(row["ticket"], []).append(row)
    author = options["author_display"] or options["author"]
    for row in weekly_rows:
        row["author"] = author
        items = by_ticket.get(row["ticket"], [])
        total = 0.0
        progress = ""
        remark = ""
        for item in items:
            total += _number(item.get("hours"))
            if str(item.get("progress") or "").strip():
                progress = item["progress"]
            if str(item.get("remark") or "").strip():
                remark = item["remark"]
        if total:
            row["hours"] = (
                str(int(total))
                if float(total).is_integer()
                else str(round(total, 2))
            )
        else:
            row["hours"] = ""
        row["progress"] = progress
        row["remark"] = remark
    return weekly_rows


def _clean_text(value):
    return "" if value is None else str(value).strip()


def _pick_write_date(row, mode, week_start, aggregate_mode="daily"):
    if (aggregate_mode or "daily") == "daily":
        day = row.get("day") or row.get("date_range")
        text = str(day or "")[:10]
        try:
            return datetime.strptime(text, "%Y-%m-%d").date()
        except ValueError:
            pass
    mode = (mode or "last").strip().lower()
    if mode == "range":
        return row.get("date_range") or ""
    if mode == "week_start":
        return week_start
    key = "first_at" if mode == "first" else "last_at"
    text = str(row.get(key) or "")[:10]
    try:
        return datetime.strptime(text, "%Y-%m-%d").date()
    except ValueError:
        return row.get("date_range") or ""


def merge_rows(ui_rows, aggregated, aggregate_mode="daily"):
    by_key = {
        row.get("row_key") or row["ticket"]: row for row in aggregated
    }
    merged = []
    for item in ui_rows:
        if not isinstance(item, dict):
            continue
        ticket = _clean_text(item.get("ticket")) or git_source.UNSORTED_TICKET
        key = item.get("row_key")
        if not key:
            if (aggregate_mode or "daily") == "daily" and item.get("date_range"):
                key = "{}|{}".format(str(item["date_range"])[:10], ticket)
            else:
                key = ticket
        base = by_key.get(key) or by_key.get(ticket)
        if base is None:
            base = {
                "ticket": ticket,
                "ticket_raw": "",
                "description": _clean_text(item.get("description")),
                "repos": _clean_text(item.get("repos")),
                "branches": _clean_text(item.get("branches")),
                "date_range": _clean_text(item.get("date_range")),
                "commit_count": int(item.get("commit_count") or 0),
            }
        row = dict(base)
        row["description"] = (
            _clean_text(item.get("description")) or row.get("description", "")
        )
        row["hours"] = _clean_text(item.get("hours"))
        row["progress"] = _clean_text(item.get("progress"))
        row["remark"] = _clean_text(item.get("remark"))
        merged.append(row)
    return merged


def _safe_name(name):
    base = Path(str(name or "")).name
    base = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", base).strip()
    if len(base) > 120:
        path = Path(base)
        base = path.stem[:100] + path.suffix
    return base or "workbook.xlsx"


def _snapshot_path(run_id, url=""):
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    name = _safe_name(excel_writer.workbook_name_from_url(url))
    return SNAPSHOT_DIR / "run_{}_{}_{}".format(run_id, stamp, name)


def start_fill(week_start_value, rows, dry_run, author=None, range_end=None):
    if not isinstance(rows, list) or not rows:
        raise ValidationError("没有可写入的汇总行")
    options = load_options()
    if not options["sharepoint_url"]:
        raise ValidationError("请先在设置里填写 SharePoint 文档链接")
    if not options["repos"]:
        raise ValidationError("请先在设置里配置 Git 仓库路径")

    options["author"] = (
        options["author"] if author is None else str(author).strip()
    )

    if not _fill_lock.acquire(blocking=False):
        raise FillBusy("已有一个填充任务在运行，请等待完成")

    try:
        week_start = _parse_anchor(week_start_value)
        if range_end:
            week_end = _parse_anchor(range_end)
            if week_start > week_end:
                raise ValidationError("开始日期不能晚于结束日期")
        else:
            week_start, week_end = git_source.week_bounds(
                week_start, options["week_start_day"]
            )
        with SessionLocal() as session:
            run = FillRun(
                week_start=week_start,
                week_end=week_end,
                author=options["author"],
                repos=json.dumps(options["repos"], ensure_ascii=False),
                mode="live" if not dry_run else "dry_run",
                status="pending",
                target_url=options["sharepoint_url"],
                rows_count=len(rows),
            )
            session.add(run)
            session.commit()
            run_id = run.id
            log_path = LOG_DIR / "run_{}.log".format(run_id)
            run.log_path = str(log_path)
            session.commit()

        thread = threading.Thread(
            target=_execute_fill,
            args=(run_id, rows, bool(dry_run), options, week_start, week_end),
            daemon=True,
            name="fill-{}".format(run_id),
        )
        thread.start()
        return run_id
    except Exception:
        _fill_lock.release()
        raise


def _execute_fill(run_id, ui_rows, dry_run, options, week_start, week_end):
    session = SessionLocal()
    logger = RunLogger(run_id, LOG_DIR / "run_{}.log".format(run_id))
    try:
        run = session.get(FillRun, run_id)
        if run is None:
            return
        run.status = "running"
        run.started_at = datetime.now()
        session.commit()

        logger.info(
            "开始{}任务 #{}：{} ~ {}，作者 {}".format(
                "演练" if dry_run else "填充",
                run_id,
                week_start,
                week_end,
                options["author"] or "全部",
            )
        )
        logger.info("仓库: {}".format("；".join(options["repos"])))

        commits = git_source.load_commits(
            options["repos"],
            options["author"],
            week_start,
            week_end,
            ignore_merges=options["ignore_merge_commits"],
        )
        logger.info("采集到 {} 条 Git 提交".format(len(commits)))
        for commit in commits:
            session.add(
                FillCommit(
                    run_id=run_id,
                    commit_hash=commit.commit_hash,
                    author=commit.author,
                    committed_at=commit.committed_at,
                    repo=commit.repo,
                    branch=commit.branch,
                    subject=commit.subject,
                    ticket=git_source.extract_ticket(commit.subject)
                    or git_source.extract_ticket(commit.body),
                )
            )
        run.commits_count = len(commits)

        target_mode = options["target_mode"]
        if target_mode == "weekly":
            aggregate_mode = "ticket"
        elif target_mode == "both":
            aggregate_mode = "daily"
        else:
            aggregate_mode = options["aggregate_mode"]
        aggregated = git_source.aggregate_commits(
            commits, aggregate_mode, options["max_description_chars"]
        )
        merged = merge_rows(ui_rows, aggregated, aggregate_mode)
        if not merged:
            raise RuntimeError("没有可写入的汇总行")
        for row in merged:
            row["author"] = options["author_display"] or options["author"]
            row["write_date"] = _pick_write_date(
                row, options["date_mode"], week_start, aggregate_mode
            )
        for seq, row in enumerate(merged, 1):
            session.add(
                FillRow(
                    run_id=run_id,
                    seq=seq,
                    ticket=row["ticket"],
                    description=row["description"],
                    repos=row["repos"],
                    branches=row["branches"],
                    date_range=row["date_range"],
                    commit_count=row["commit_count"],
                    hours=row["hours"],
                    progress=row["progress"],
                    remark=row["remark"],
                )
            )
        run.rows_count = len(merged)
        session.commit()
        logger.info(
            "汇总为 {} 行（{}）".format(
                len(merged),
                "按 Ticket + 日期"
                if aggregate_mode == "daily"
                else "按 Ticket",
            )
        )

        snapshot_path = _snapshot_path(run_id, options["sharepoint_url"])
        base_options = dict(options)
        base_options["target_workbook"] = excel_writer.workbook_name_from_url(
            options["sharepoint_url"]
        )
        base_options["weekly_template_sheet"] = options["weekly_template_sheet"]
        base_options["weekly_sheet_prefix"] = options["weekly_sheet_prefix"]
        base_options["weekly_overwrite"] = options["weekly_overwrite"]
        base_options["range_start"] = week_start.isoformat()

        warnings = []
        weekly_rows = None
        if target_mode in ("weekly", "both"):
            if target_mode == "weekly":
                weekly_rows = merged
            else:
                weekly_rows = git_source.aggregate_commits(
                    commits, "ticket", options["max_description_chars"]
                )
                if options["ai_enabled"]:
                    ai_rewrite.rewrite_rows(weekly_rows, options, warnings)
                _enrich_weekly_rows(merged, weekly_rows, options)

        result = None
        weekly_result = None
        if target_mode in ("existing", "both"):
            existing_options = dict(base_options)
            existing_options["target_mode"] = "existing"
            existing_options["snapshot_path"] = str(snapshot_path)
            result = excel_writer.write_rows(
                options["sharepoint_url"],
                merged,
                existing_options,
                logger,
                dry_run,
            )
        if target_mode in ("weekly", "both"):
            weekly_options = dict(base_options)
            weekly_options["target_mode"] = "weekly"
            weekly_options["snapshot_path"] = (
                "" if target_mode == "both" else str(snapshot_path)
            )
            weekly_result = excel_writer.write_rows(
                options["sharepoint_url"],
                weekly_rows,
                weekly_options,
                logger,
                dry_run,
            )
            if result is None:
                result = weekly_result

        if result is None:
            raise RuntimeError("没有可写入的目标")

        run.sheet_name = result.plan.sheet_name
        if target_mode == "both" and weekly_result is not None:
            run.sheet_name = "{} + {}".format(
                result.plan.sheet_name, weekly_result.plan.sheet_name
            )[:200]
        run.header_row = result.plan.header_row
        run.start_row = result.plan.start_row
        run.end_row = result.plan.end_row
        run.snapshot_path = result.snapshot_path
        run.audit_path = str(AUDIT_DIR / "run_{}.xlsx".format(run_id))

        if target_mode == "both" and weekly_rows is not None:
            for index, row in enumerate(weekly_rows, 1):
                row["excel_row"] = (weekly_result.plan.start_row or 0) + index - 1

        if target_mode == "both":
            run.message = (
                "演练完成：日报 {} 行 + 周报汇总 {} 行（未保存）".format(
                    len(merged), len(weekly_rows or [])
                )
                if dry_run
                else "日报已写入 {} 行；周报汇总「{}」已写入 {} 行".format(
                    len(merged),
                    weekly_result.plan.sheet_name,
                    len(weekly_rows or []),
                )
            )
        elif target_mode == "weekly":
            run.message = (
                "演练完成：周报汇总 {} 行（未保存）".format(len(weekly_rows or []))
                if dry_run
                else "周报汇总「{}」已写入 {} 行".format(
                    result.plan.sheet_name, len(weekly_rows or [])
                )
            )
        else:
            run.message = (
                "演练完成，未写入 SharePoint"
                if dry_run
                else "已写入并保存到 SharePoint"
            )

        if result.warnings:
            warnings.extend(result.warnings)
        if weekly_result is not None and weekly_result is not result:
            warnings.extend(weekly_result.warnings)
        if warnings:
            run.message += "；" + "；".join(warnings)

        db_rows = (
            session.query(FillRow)
            .filter(FillRow.run_id == run_id)
            .order_by(FillRow.seq)
            .all()
        )
        for db_row in db_rows:
            db_row.excel_row = result.plan.start_row + db_row.seq - 1
            db_row.written = not dry_run
        session.commit()

        run.status = "success"
        run.finished_at = datetime.now()
        session.commit()
        logger.info("任务完成")

        try:
            export_audit(
                run,
                merged,
                commits,
                logger.text(),
                Path(run.audit_path),
                weekly_rows=weekly_rows if target_mode == "both" else None,
            )
            logger.info("已导出审计表: {}".format(run.audit_path))
        except Exception:
            logger.exception("审计表导出失败（不影响填充结果）")
            run.message += "；审计表导出失败"
            session.commit()
    except Exception as error:
        logger.exception("任务失败")
        try:
            run = session.get(FillRun, run_id)
            if run is not None:
                run.status = "failed"
                run.message = str(error)
                run.finished_at = datetime.now()
                session.commit()
        except Exception:
            session.rollback()
    finally:
        session.close()
        _fill_lock.release()


def list_runs(limit=50, week_start=None, week_end=None, mode=None, status=None):
    with SessionLocal() as session:
        query = session.query(FillRun)
        if week_start:
            query = query.filter(FillRun.week_start == _parse_anchor(week_start))
        if week_end:
            query = query.filter(FillRun.week_end == _parse_anchor(week_end))
        if mode in ("dry_run", "live"):
            query = query.filter(FillRun.mode == mode)
        if status in ("pending", "running", "success", "failed"):
            query = query.filter(FillRun.status == status)
        runs = (
            query.order_by(FillRun.id.desc())
            .limit(max(1, min(int(limit or 50), 200)))
            .all()
        )
        return [run.to_dict() for run in runs]


def run_detail(run_id):
    with SessionLocal() as session:
        run = session.get(FillRun, run_id)
        if run is None:
            return None
        data = run.to_dict(include_children=True)
        log_text = ""
        if run.log_path and Path(run.log_path).exists():
            try:
                log_text = Path(run.log_path).read_text(encoding="utf-8")
            except OSError:
                log_text = ""
        data["log"] = log_text
        return data


def retry_run(run_id):
    with SessionLocal() as session:
        run = session.get(FillRun, run_id)
        if run is None:
            raise ValidationError("填充记录不存在")
        if run.status == "running":
            raise ValidationError("该任务仍在运行")
        rows = [
            {
                "ticket": row.ticket,
                "description": row.description,
                "repos": row.repos,
                "branches": row.branches,
                "date_range": row.date_range,
                "commit_count": row.commit_count,
                "hours": row.hours,
                "progress": row.progress,
                "remark": row.remark,
            }
            for row in sorted(run.rows, key=lambda item: item.seq)
        ]
        week_start = run.week_start
        week_end = run.week_end
        dry_run = run.mode == "dry_run"
        run_author = run.author
    return start_fill(
        week_start.isoformat() if week_start else None,
        rows,
        dry_run,
        author=run_author,
        range_end=week_end.isoformat() if week_end else None,
    )


def rewrite_ai(rows):
    options = load_options()
    if not options["ai_enabled"]:
        raise ValidationError("请先在设置里启用 AI 改写")
    cleaned = []
    for item in rows or []:
        if not isinstance(item, dict):
            continue
        cleaned.append(dict(item))
    if not cleaned:
        raise ValidationError("没有可改写的行")
    warnings = []
    ai_rewrite.rewrite_rows(cleaned, options, warnings)
    return {"rows": cleaned, "warnings": warnings}


def test_ai():
    options = load_options()
    if not (
        options["ai_base_url"]
        and options["ai_api_key"]
        and options["ai_model"]
    ):
        raise ValidationError("请先填写 AI Base URL、API Key 和模型")
    return ai_rewrite.test_connection(options)


def _style_header(sheet, row_number, headers):
    font = Font(bold=True, color="1F3A5F")
    fill = PatternFill("solid", fgColor="DCE8FA")
    center = Alignment(horizontal="center", vertical="center")
    thin = Side(style="thin", color="B9CBE6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for column, title in enumerate(headers, 1):
        cell = sheet.cell(row=row_number, column=column, value=title)
        cell.font = font
        cell.fill = fill
        cell.alignment = center
        cell.border = border


def _repos_of(run):
    try:
        return json.loads(run.repos) if run.repos else []
    except (ValueError, TypeError):
        return [run.repos] if run.repos else []


def export_audit(run, rows, commits, log_text, output_path, weekly_rows=None):
    workbook = openpyxl.Workbook()
    info = workbook.active
    info.title = "运行信息"
    info_rows = [
        ("运行编号", run.id),
        ("周区间", "{} ~ {}".format(run.week_start, run.week_end)),
        ("作者", run.author or "全部"),
        ("仓库", "；".join(_repos_of(run))),
        ("模式", "演练（未写入）" if run.mode == "dry_run" else "正式写入"),
        ("状态", run.status),
        ("目标文档", run.target_url),
        ("工作表", run.sheet_name),
        ("表头行", run.header_row),
        ("写入区间", "{} ~ {}".format(run.start_row, run.end_row)),
        ("写入行数", run.rows_count),
        ("提交数", run.commits_count),
        ("写入前快照", run.snapshot_path),
        ("开始时间", run.started_at),
        ("结束时间", run.finished_at),
        ("说明", run.message),
    ]
    for index, (key, value) in enumerate(info_rows, 1):
        info.cell(row=index, column=1, value=key).font = Font(bold=True)
        info.cell(row=index, column=2, value="" if value is None else str(value))
    info.column_dimensions["A"].width = 16
    info.column_dimensions["B"].width = 90
    for row in info.iter_rows(min_row=1, max_row=len(info_rows), min_col=1, max_col=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    summary = workbook.create_sheet("填充记录")
    headers = [
        "序号",
        "Ticket",
        "任务描述",
        "仓库",
        "分支",
        "日期",
        "提交数",
        "工时(h)",
        "进度(%)",
        "备注",
        "Excel行号",
    ]
    _style_header(summary, 1, headers)
    for offset, row in enumerate(rows, 2):
        values = [
            offset - 1,
            row["ticket"],
            row["description"],
            row["repos"],
            row["branches"],
            row["date_range"],
            row["commit_count"],
            row["hours"],
            row["progress"],
            row["remark"],
            (run.start_row or 0) + offset - 2,
        ]
        for column, value in enumerate(values, 1):
            cell = summary.cell(row=offset, column=column, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(
                left=Side(style="thin", color="D9E2EF"),
                right=Side(style="thin", color="D9E2EF"),
                top=Side(style="thin", color="D9E2EF"),
                bottom=Side(style="thin", color="D9E2EF"),
            )
    for column, width in enumerate([7, 18, 60, 22, 18, 24, 9, 9, 9, 24, 10], 1):
        summary.column_dimensions[get_column_letter(column)].width = width
    summary.freeze_panes = "A2"

    if weekly_rows:
        weekly = workbook.create_sheet("周报汇总")
        weekly_headers = [
            "序号",
            "执行人",
            "完成事项",
            "Ticket/需求编号",
            "工时",
            "状态",
            "进度",
            "Excel行号",
        ]
        _style_header(weekly, 1, weekly_headers)
        for offset, row in enumerate(weekly_rows, 2):
            values = [
                offset - 1,
                row.get("author", ""),
                row.get("description", ""),
                row.get("ticket", ""),
                row.get("hours", ""),
                row.get("remark", ""),
                row.get("progress", ""),
                row.get("excel_row", ""),
            ]
            for column, value in enumerate(values, 1):
                cell = weekly.cell(row=offset, column=column, value=value)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = Border(
                    left=Side(style="thin", color="D9E2EF"),
                    right=Side(style="thin", color="D9E2EF"),
                    top=Side(style="thin", color="D9E2EF"),
                    bottom=Side(style="thin", color="D9E2EF"),
                )
        for column, width in enumerate([7, 14, 70, 18, 9, 14, 9, 10], 1):
            weekly.column_dimensions[get_column_letter(column)].width = width
        weekly.freeze_panes = "A2"

    detail = workbook.create_sheet("Git提交明细")
    commit_headers = [
        "提交哈希",
        "作者",
        "提交时间",
        "仓库",
        "分支",
        "Ticket",
        "提交说明",
    ]
    _style_header(detail, 1, commit_headers)
    for offset, commit in enumerate(commits, 2):
        values = [
            commit.commit_hash[:12],
            commit.author,
            commit.committed_at.strftime("%Y-%m-%d %H:%M:%S"),
            commit.repo,
            commit.branch,
            git_source.extract_ticket(commit.subject)
            or git_source.extract_ticket(commit.body),
            commit.subject,
        ]
        for column, value in enumerate(values, 1):
            detail.cell(row=offset, column=column, value=value).alignment = Alignment(
                vertical="top", wrap_text=True
            )
    for column, width in enumerate([14, 14, 20, 20, 18, 16, 80], 1):
        detail.column_dimensions[get_column_letter(column)].width = width
    detail.freeze_panes = "A2"

    log_sheet = workbook.create_sheet("运行日志")
    for index, line in enumerate(log_text.splitlines(), 1):
        log_sheet.cell(row=index, column=1, value=line)
    log_sheet.column_dimensions["A"].width = 140

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def list_samples():
    samples = []
    for path in sorted(SAMPLES_DIR.glob("*")):
        if path.suffix.lower() not in (".xlsx", ".xlsm"):
            continue
        stat = path.stat()
        samples.append(
            {
                "name": path.name,
                "path": str(path),
                "size": stat.st_size,
                "modified_at": datetime.fromtimestamp(stat.st_mtime).strftime(
                    "%Y-%m-%d %H:%M:%S"
                ),
            }
        )
    return samples


def save_sample(file: FileStorage):
    if file is None or not file.filename:
        raise ValidationError("没有选择文件")
    name = _safe_name(file.filename)
    if Path(name).suffix.lower() not in (".xlsx", ".xlsm"):
        raise ValidationError("只支持 .xlsx / .xlsm 文件")
    target = SAMPLES_DIR / name
    file.save(str(target))
    return {"name": target.name, "path": str(target)}


def inspect_template(path_value):
    path = Path(str(path_value or "")).expanduser()
    if not path.is_absolute():
        path = SAMPLES_DIR / path
    if not path.exists():
        raise ValidationError("模板文件不存在: {}".format(path))
    if path.suffix.lower() not in (".xlsx", ".xlsm"):
        raise ValidationError("只支持 .xlsx / .xlsm 文件")
    options = load_options()
    return excel_writer.inspect_local(
        path, options["aliases"], options["header_row"]
    )


def inspect_sharepoint():
    options = load_options()
    if not options["sharepoint_url"]:
        raise ValidationError("请先在设置里填写 SharePoint 文档链接")
    sheet = (
        options["weekly_template_sheet"]
        if options["target_mode"] == "weekly"
        else options["target_sheet"]
    )
    return excel_writer.inspect_sharepoint(
        options["sharepoint_url"],
        options["aliases"],
        sheet,
        options["header_row"],
        write_mode=options["write_mode"],
        target_workbook=excel_writer.workbook_name_from_url(
            options["sharepoint_url"]
        ),
    )


def workbook_status():
    options = load_options()
    return excel_writer.workbook_status(
        options["sharepoint_url"],
        excel_writer.workbook_name_from_url(options["sharepoint_url"]),
    )
